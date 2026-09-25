"""
Resume Parser
=============
Extracts text from PDF, DOCX and TXT resume files.
"""

import os
import re


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):
    """
    Clean extracted resume text.
    """

    if not text:
        return ""

    text = text.replace("\x00", " ")

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove excessive spaces
    text = re.sub(r"[ \t]+", " ", text)

    # Remove excessive blank lines
    text = re.sub(r"\n\s*\n+", "\n\n", text)

    return text.strip()


# ============================================================
# PDF
# ============================================================

def extract_from_pdf(file_path):
    """
    Extract text from PDF.
    """

    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise ImportError(
            "PyPDF2 is not installed. Run: "
            "python -m pip install PyPDF2"
        )

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    reader = PdfReader(file_path)

    pages_text = []

    for page in reader.pages:

        try:
            page_text = page.extract_text()
        except Exception:
            page_text = ""

        if page_text:
            pages_text.append(page_text)

    text = "\n".join(pages_text)

    return clean_text(text)


# ============================================================
# DOCX
# ============================================================

def extract_from_docx(file_path):
    """
    Extract text from DOCX.
    """

    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is not installed. Run: "
            "python -m pip install python-docx"
        )

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    document = Document(file_path)

    paragraphs = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                paragraphs.append(
                    " | ".join(row_text)
                )

    return clean_text(
        "\n".join(paragraphs)
    )


# ============================================================
# TXT
# ============================================================

def extract_from_txt(file_path):
    """
    Extract text from TXT.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:

        text = file.read()

    return clean_text(text)


# ============================================================
# GENERAL RESUME EXTRACTOR
# ============================================================

def extract_resume_text(file_path):
    """
    Automatically determine file type
    and extract text.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    if extension == ".docx":
        return extract_from_docx(file_path)

    if extension == ".txt":
        return extract_from_txt(file_path)

    raise ValueError(
        f"Unsupported file format: {extension}"
    )


# ============================================================
# COMPATIBILITY FUNCTIONS
# ============================================================

def extract_pdf_text(file_path):
    return extract_from_pdf(file_path)


def extract_docx_text(file_path):
    return extract_from_docx(file_path)


def extract_txt_text(file_path):
    return extract_from_txt(file_path)